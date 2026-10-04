#!/usr/bin/env python3
"""Offline Chapter 9 fictional cost example. No network or API calls.

Write formulas with Python-computed cached values, then inspect both
with openpyxl. The restricted evaluator is not a general Excel engine.
"""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import platform
import re
import tempfile
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

import openpyxl
import xlsxwriter
from openpyxl.utils.cell import get_column_letter, range_boundaries


SOURCE_DATE = "2026-10-03"
INPUTS = {"B21": Decimal("48000.00"), "B22": Decimal("36000.00"),
          "B23": Decimal("15000.00")}
LABELS = {"B5": "Recurring annual saving", "B6": "First-year proposed cost",
          "B7": "First-year additional cost", "B8": "Five-year current cost",
          "B9": "Five-year proposed cost", "B10": "Five-year saving"}
FORMULAS = {"B5": "=B21-B22", "B6": "=E14", "B7": "=E14-B14",
            "B8": "=SUM(B14:B18)", "B9": "=SUM(E14:E18)", "B10": "=B8-B9"}
for row in range(14, 19):
    FORMULAS.update({f"B{row}": "=$B$21", f"C{row}": "=$B$22",
                     f"E{row}": f"=SUM(C{row}:D{row})", f"F{row}": f"=B{row}-E{row}",
                     f"G{row}": f"=SUM(F$14:F{row})"})
FORMULAS["D14"] = "=$B$23"
CONSTANTS = {f"D{row}": Decimal("0") for row in range(15, 19)}


def require(condition: bool, message: str) -> None:
    """Keep checks active even when Python runs with -O."""
    if not condition:
        raise AssertionError(message)


def money(value: Decimal) -> Decimal:
    if not isinstance(value, Decimal) or not value.is_finite():
        raise ValueError("A finite Decimal is required; missing is not zero")
    if value < 0 or value != value.quantize(Decimal("0.01")):
        raise ValueError("This example requires nonnegative USD inputs in cents")
    return value


def excel_number(value: Decimal) -> int | float:
    """Keep whole USD exact; convert cents only for XLSX numeric storage."""
    return int(value) if value == value.to_integral_value() else float(value)


def calculate(inputs: dict[str, Decimal], formulas=FORMULAS) -> dict[str, Decimal]:
    """Execute only the explicit arithmetic used here, with exact decimals."""
    values = {**CONSTANTS, **{key: money(value) for key, value in inputs.items()}}
    visiting: set[str] = set()

    def lookup(address: str) -> Decimal:
        if address in values:
            return values[address]
        if address in visiting:
            raise ValueError(f"Circular reference: {address}")
        if address not in formulas:
            raise ValueError(f"Missing cell: {address}")
        visiting.add(address)
        expression = formulas[address].removeprefix("=").replace("$", "")
        range_sum = re.fullmatch(r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)", expression)
        if range_sum:
            left, top, right, bottom = range_boundaries(":".join(range_sum.groups()))
            result = sum((lookup(f"{get_column_letter(col)}{row}")
                          for row in range(top, bottom + 1)
                          for col in range(left, right + 1)), Decimal("0"))
        else:
            expression = re.sub(r"\b[A-Z]+\d+\b", lambda match: f'C("{match[0]}")', expression)

            def evaluate(node: ast.AST) -> Decimal:
                if isinstance(node, ast.Expression):
                    return evaluate(node.body)
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
                    if node.func.id == "C" and len(node.args) == 1 and not node.keywords:
                        arg = node.args[0]
                        if isinstance(arg, ast.Constant) and isinstance(arg.value, str):
                            return lookup(arg.value)
                if isinstance(node, ast.BinOp):
                    left_value, right_value = evaluate(node.left), evaluate(node.right)
                    if isinstance(node.op, ast.Add):
                        return left_value + right_value
                    if isinstance(node.op, ast.Sub):
                        return left_value - right_value
                    if isinstance(node.op, ast.Mult):
                        return left_value * right_value
                raise ValueError(f"Unsupported expression in {address}")

            result = evaluate(ast.parse(expression, mode="eval"))
        visiting.remove(address)
        values[address] = result
        return result

    return {address: lookup(address) for address in formulas}


def check_calculations() -> list[dict]:
    checks = []

    def check(name, observed, expected):
        require(observed == expected, f"{name}: {observed} != {expected}")
        checks.append({"check": name, "expected": str(expected),
                       "observed": str(observed), "outcome": "passed"})

    values = calculate(INPUTS)
    expected = {"B5": "12000", "B6": "51000", "B7": "3000",
                "B8": "240000", "B9": "195000", "B10": "45000"}
    for address, value in expected.items():
        check(LABELS[address], values[address], Decimal(value))
    check("Five-year saving equals sum of annual savings", values["B10"],
          sum((values[f"F{row}"] for row in range(14, 19)), Decimal("0")))
    check("First-year extra cost reconciles to implementation less annual saving",
          values["B7"], INPUTS["B23"] - values["B5"])
    for row, total in zip(range(14, 19), ("-3000", "9000", "21000", "33000", "45000")):
        check(f"Year {row - 13} cumulative saving", values[f"G{row}"], Decimal(total))
    zero_install = calculate({**INPUTS, "B23": Decimal("0")})
    check("Zero implementation cost: first-year saving", -zero_install["B7"], Decimal("12000"))
    high_install = calculate({**INPUTS, "B23": Decimal("75000")})
    check("Higher implementation cost reverses five-year saving", high_install["B10"], Decimal("-15000"))
    cents = calculate({**INPUTS, "B22": Decimal("36000.01")})
    check("Decimal cents retained over five years", cents["B9"], Decimal("195000.05"))
    try:
        calculate({**INPUTS, "B22": None})
    except ValueError:
        checks.append({"check": "Missing input rejected", "outcome": "passed"})
    else:
        raise AssertionError("A missing cost was accepted")
    return checks


def create_workbook(path: Path, values: dict[str, Decimal]) -> None:
    """Store formulas and executed Python values; do not claim Excel ran."""
    with xlsxwriter.Workbook(path) as book:
        book.set_properties({"title": "Project cost comparison", "author": "World Agrees reader resources",
                             "created": datetime.fromisoformat(SOURCE_DATE),
                             "comments": "Fictional Chapter 9 example; Python-computed formula caches."})
        book.set_calc_mode("auto")
        sheet = book.add_worksheet("Costs")
        sheet.hide_gridlines(2)
        sheet.set_tab_color("#243B53")
        sheet.set_column("A:A", 36)
        sheet.set_column("B:C", 18)
        sheet.set_column("D:D", 20)
        sheet.set_column("E:E", 18)
        sheet.set_column("F:G", 22)
        sheet.set_default_row(21)
        base = {"font_name": "Arial", "font_size": 11, "valign": "vcenter"}
        money_format = '$#,##0.00;($#,##0.00);"—"'
        title = book.add_format({**base, "bold": True, "font_size": 15})
        text = book.add_format(base)
        centered_text = book.add_format({**base, "align": "center"})
        note = book.add_format({**base, "font_size": 10, "italic": True})
        currency = book.add_format({**base, "num_format": money_format})
        input_format = book.add_format({**base, "font_color": "#185ABD", "num_format": money_format})
        header = book.add_format({**base, "bold": True, "font_color": "#FFFFFF",
                                  "bg_color": "#243B53", "align": "center"})
        date_format = book.add_format({**base, "num_format": "yyyy-mm-dd"})
        sheet.write("A2", "Project cost comparison", title)
        sheet.write("A3", "Hypothetical example. All monetary values are USD.", note)
        for address, label in LABELS.items():
            sheet.write(f"A{address[1:]}", label, text)
        sheet.write_row("A13", ["Year", "Current cost", "Proposed annual", "Implementation",
                                "Proposed total", "Saving", "Cumulative saving"], header)
        for row in range(14, 19):
            sheet.write_number(f"A{row}", row - 13, text)
        for address, value in CONSTANTS.items():
            sheet.write_number(address, int(value), currency)
        sheet.write_row("A20", ["Input", "Amount (USD)", "Period", "Source ID", "Source date", "Status"], header)
        for row, label, period, source in [(21, "Current annual cost", "Annual", "H1"),
                                          (22, "Proposed annual cost", "Annual", "H2"),
                                          (23, "Implementation cost", "First year only", "H3")]:
            sheet.write(f"A{row}", label, text)
            sheet.write_number(f"B{row}", excel_number(INPUTS[f"B{row}"]), input_format)
            sheet.write(f"C{row}", period, centered_text)
            sheet.write(f"D{row}", source, text)
            sheet.write_datetime(f"E{row}", datetime.fromisoformat(SOURCE_DATE), date_format)
            sheet.write(f"F{row}", "Fictional example", centered_text)
        for address, formula in FORMULAS.items():
            sheet.write_formula(address, formula, currency, excel_number(values[address]))
        sheet.write("A25", "Source: Chapter 9 fictional inputs, version 2026-10-03; not quotes or observed costs.", note)
        sheet.write("A26", "Assumptions: unchanged annual rates; implementation in year 1; no other costs, inflation, or discounting.", note)
        sheet.write("A27", "Caches computed by Python for these inputs. Recalculate after edits; native Excel/LibreOffice not tested here.", note)
        sheet.set_landscape()
        sheet.fit_to_pages(1, 1)
        sheet.print_area("A1:G27")


def verify_saved_workbook(path: Path) -> list[dict]:
    formula_book = openpyxl.load_workbook(path, data_only=False)
    value_book = openpyxl.load_workbook(path, data_only=True)
    formula_sheet, value_sheet = formula_book["Costs"], value_book["Costs"]
    saved_inputs = {address: Decimal(str(formula_sheet[address].value)) for address in INPUTS}
    require(saved_inputs == INPUTS, "Saved inputs differ from the example sources")
    saved_formulas = {address: formula_sheet[address].value for address in FORMULAS}
    require(saved_formulas == FORMULAS, "Saved formulas differ from the declared formulas")
    recomputed = calculate(saved_inputs, saved_formulas)
    checks = []
    for address, expected in recomputed.items():
        cached = value_sheet[address].value
        require(cached is not None, f"Missing formula cache: {address}")
        require(Decimal(str(cached)) == expected, f"Formula/cache mismatch: {address}")
        checks.append({"cell": f"Costs!{address}", "formula": saved_formulas[address],
                       "cached_value": str(cached), "recomputed_value": str(expected), "outcome": "passed"})
    formula_book.close()
    value_book.close()
    return checks


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path,
                        help="New/empty directory outside reader resources; defaults to a temporary directory")
    args = parser.parse_args()
    output = args.output_dir.resolve() if args.output_dir else Path(tempfile.mkdtemp(prefix="world-agrees-project-costs-"))
    source_root = Path(__file__).resolve().parents[2]
    if output == source_root or source_root in output.parents:
        parser.error("Outputs must be outside the reader-resources source tree")
    if output.exists() and any(output.iterdir()):
        parser.error("Output directory must be empty; existing files will not be overwritten")
    output.mkdir(parents=True, exist_ok=True)
    checks = check_calculations()
    values = calculate(INPUTS)
    path = output / "project-costs.xlsx"
    create_workbook(path, values)
    formula_checks = verify_saved_workbook(path)
    with (output / "results.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["cell", "metric", "amount", "currency"])
        writer.writerows([address, label, f"{values[address]:.2f}", "USD"] for address, label in LABELS.items())
    report = {
        "status": "passed", "executed_at": datetime.now(timezone.utc).isoformat(),
        "source": "Chapter 9 hypothetical example; not observed costs", "source_version_date": SOURCE_DATE,
        "inputs": {key: str(value) for key, value in INPUTS.items()},
        "calculation_engine": "Python Decimal restricted evaluation of this script's explicit formulas",
        "cache_origin": "Executed Python calculation supplied to xlsxwriter; inspected with openpyxl",
        "runtime": f"Python {platform.python_version()}",
        "libraries": {"openpyxl": openpyxl.__version__, "xlsxwriter": xlsxwriter.__version__},
        "native_spreadsheet_application_recalculation": "not performed",
        "assumptions": ["five undiscounted years", "unchanged annual costs", "implementation in year one",
                        "no other costs, taxes, inflation, financing, or savings assumed"],
        "calculation_checks": checks, "saved_formula_cache_check_count": len(formula_checks),
        "saved_formula_cache_outcome": "all passed",
        "summary_cell_checks": [check for check in formula_checks if check["cell"].split("!")[1] in LABELS],
        "workbook_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }
    (output / "verification.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output_dir": str(output), "status": report["status"],
                      "calculation_checks": len(checks), "formula_cache_checks": len(formula_checks),
                      "five_year_saving_usd": str(values["B10"]),
                      "native_application_recalculation": "not performed"}, indent=2))


if __name__ == "__main__":
    main()
